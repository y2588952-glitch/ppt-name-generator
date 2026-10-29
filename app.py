# -*- coding: utf-8 -*-
"""PPT 台卡批量生成工具 —— Flask 后端入口。

所有文件均在本地处理，不依赖任何云端接口。
上传文件不会被提交到 GitHub（见 .gitignore）。
"""

import os
import re
import shutil
import time
import uuid

from flask import Flask, jsonify, render_template, request, send_from_directory
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

import ppt_processor
from ppt_processor import PptError

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMP_DIR = os.path.join(BASE_DIR, "temp")
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")
STATIC_DIR = os.path.join(BASE_DIR, "static")

MAX_PPT_SIZE = 50 * 1024 * 1024      # 50 MB
MAX_NAMES_SIZE = 20 * 1024 * 1024    # 20 MB
TASK_TTL_SECONDS = 24 * 3600         # 临时目录保留 24 小时

ALLOWED_PPT_EXT = {".pptx"}
ALLOWED_NAMES_EXT = {".txt", ".xlsx", ".xlsm"}

TASK_ID_RE = re.compile(r"^[A-Za-z0-9_-]{8,64}$")

app = Flask(__name__, template_folder=TEMPLATE_DIR, static_folder=STATIC_DIR)
app.config["MAX_CONTENT_LENGTH"] = 60 * 1024 * 1024
app.config["JSON_AS_ASCII"] = False


class ApiError(Exception):
    """带 HTTP 状态码的业务错误。"""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message = message
        self.status = status


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------
def _sweep_old_tasks() -> None:
    """删除超过 TASK_TTL_SECONDS 的临时目录。"""
    if not os.path.isdir(TEMP_DIR):
        return
    now = time.time()
    for name in os.listdir(TEMP_DIR):
        full = os.path.join(TEMP_DIR, name)
        try:
            if os.path.isdir(full) and now - os.path.getmtime(full) > TASK_TTL_SECONDS:
                shutil.rmtree(full, ignore_errors=True)
        except OSError:
            pass


def _task_dir(task_id: str) -> str:
    if not TASK_ID_RE.match(task_id or ""):
        raise ApiError("无效的任务标识。")
    d = os.path.join(TEMP_DIR, task_id)
    os.makedirs(d, exist_ok=True)
    return d


def _check_ext(filename: str, allowed: set) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in allowed:
        raise ApiError("不支持的文件类型：" + (ext or "未知") + "。")
    return ext


def _original_template_name(task_dir: str) -> str:
    name_file = os.path.join(task_dir, "original_name.txt")
    if os.path.exists(name_file):
        try:
            with open(name_file, "r", encoding="utf-8") as f:
                name = f.read().strip()
                if name:
                    return name
        except OSError:
            pass
    return "台卡"


def _read_text_file(path: str) -> str:
    with open(path, "rb") as f:
        data = f.read()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_name_text(text: str, trim: bool = True) -> list:
    names = []
    for line in str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        name = line.strip() if trim else line.rstrip("\n")
        if name:
            names.append(name)
    return names


# ---------------------------------------------------------------------------
# 错误处理
# ---------------------------------------------------------------------------
@app.errorhandler(ApiError)
def _handle_api_error(err: ApiError):
    return jsonify(ok=False, error=err.message), err.status


@app.errorhandler(413)
def _handle_too_large(_err):
    return jsonify(ok=False, error="上传文件过大，请减小文件体积。"), 413


@app.errorhandler(Exception)
def _handle_generic(err):
    app.logger.exception("unhandled error")
    return jsonify(ok=False, error="服务器处理失败：" + str(err)), 500


# ---------------------------------------------------------------------------
# 页面
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ---------------------------------------------------------------------------
# 任务
# ---------------------------------------------------------------------------
@app.route("/api/new_task", methods=["POST"])
def new_task():
    _sweep_old_tasks()
    task_id = uuid.uuid4().hex[:16]
    os.makedirs(os.path.join(TEMP_DIR, task_id), exist_ok=True)
    return jsonify(ok=True, task_id=task_id)


# ---------------------------------------------------------------------------
# 上传并分析 PPT
# ---------------------------------------------------------------------------
@app.route("/api/upload_ppt", methods=["POST"])
def upload_ppt():
    task_id = request.form.get("task_id", "")
    task_dir = _task_dir(task_id)

    file = request.files.get("file")
    if file is None or not file.filename:
        raise ApiError("请选择要上传的 PPTX 文件。")

    _check_ext(file.filename, ALLOWED_PPT_EXT)

    # 清空该任务旧的模板 / 结果，避免串用
    for old in ("template.pptx", "output.pptx"):
        p = os.path.join(task_dir, old)
        if os.path.exists(p):
            os.remove(p)

    safe = secure_filename(file.filename) or "template.pptx"
    template_path = os.path.join(task_dir, "template.pptx")
    file.save(template_path)
    with open(os.path.join(task_dir, "original_name.txt"), "w", encoding="utf-8") as f:
        f.write(file.filename or "台卡")

    if os.path.getsize(template_path) > MAX_PPT_SIZE:
        os.remove(template_path)
        raise ApiError("PPT 文件过大（超过 50MB）。")

    try:
        analysis = ppt_processor.analyze_pptx(template_path)
    except Exception:
        os.remove(template_path)
        raise ApiError("无法解析该 PPTX 文件，文件可能已损坏或格式不受支持。")

    analysis["filename"] = file.filename
    return jsonify(ok=True, task_id=task_id, analysis=analysis)


# ---------------------------------------------------------------------------
# 导入姓名
# ---------------------------------------------------------------------------
@app.route("/api/import_names", methods=["POST"])
def import_names():
    task_id = request.form.get("task_id", "")
    task_dir = _task_dir(task_id)

    file = request.files.get("file")
    pasted_text = request.form.get("text", "")

    if file is not None and file.filename:
        ext = _check_ext(file.filename, ALLOWED_NAMES_EXT)
        safe = secure_filename(file.filename) or "names_input"
        name_path = os.path.join(task_dir, "names_input" + ext)
        file.save(name_path)
        if os.path.getsize(name_path) > MAX_NAMES_SIZE:
            os.remove(name_path)
            raise ApiError("名单文件过大（超过 20MB）。")

        if ext == ".txt":
            try:
                text = _read_text_file(name_path)
            except Exception:
                raise ApiError("无法读取 TXT 文件。")
            names = _parse_name_text(text)
            return jsonify(ok=True, type="names", names=names, count=len(names))

        # Excel
        try:
            wb = load_workbook(name_path, data_only=True, keep_vba=False)
        except Exception:
            raise ApiError("无法读取 Excel 文件，文件可能已损坏。")

        sheets_preview = {}
        for ws in wb.worksheets:
            rows = []
            max_col = min(ws.max_column or 1, 20)
            max_row = min(ws.max_row or 1, 8)
            for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
                rows.append(["" if c.value is None else str(c.value) for c in row])
            sheets_preview[ws.title] = rows

        is_xlsm = ext == ".xlsm"
        return jsonify(
            ok=True,
            type="excel",
            sheet_names=wb.sheetnames,
            preview=sheets_preview,
            is_xlsm=is_xlsm,
            warning=(
                "XLSM 为启用宏的工作簿。本工具只读取数据，不会执行其中的宏；"
                "原始文件不会被修改，请放心使用。"
                if is_xlsm
                else None
            ),
        )

    # 粘贴文本
    names = _parse_name_text(pasted_text)
    return jsonify(ok=True, type="names", names=names, count=len(names))


@app.route("/api/parse_excel", methods=["POST"])
def parse_excel():
    task_id = request.form.get("task_id", "")
    task_dir = _task_dir(task_id)
    sheet = request.form.get("sheet", "")
    column_raw = request.form.get("column", "")
    skip_header = request.form.get("skip_header", "0") in ("1", "true", "True")

    name_path = None
    for ext in (".xlsx", ".xlsm"):
        candidate = os.path.join(task_dir, "names_input" + ext)
        if os.path.exists(candidate):
            name_path = candidate
            break
    if name_path is None:
        raise ApiError("请先上传 Excel 名单文件。")

    try:
        column = int(column_raw)
        if column < 1:
            raise ValueError
    except (TypeError, ValueError):
        raise ApiError("请指定有效的姓名所在列。")

    try:
        wb = load_workbook(name_path, data_only=True, keep_vba=False)
        if sheet not in wb.sheetnames:
            raise ApiError("所选工作表不存在。")
        ws = wb[sheet]
        if column > (ws.max_column or 0):
            raise ApiError("所选列超出了工作表的列范围。")

        names = []
        start_row = 2 if skip_header else 1
        for row in ws.iter_rows(min_row=start_row, min_col=column, max_col=column):
            v = row[0].value
            if v is None:
                continue
            s = str(v).strip()
            if s:
                names.append(s)
    except ApiError:
        raise
    except Exception:
        raise ApiError("读取 Excel 数据失败。")

    return jsonify(ok=True, names=names, count=len(names))


# ---------------------------------------------------------------------------
# 生成 PPT
# ---------------------------------------------------------------------------
@app.route("/api/generate", methods=["POST"])
def generate():
    data = request.get_json(silent=True) or {}
    task_id = data.get("task_id", "")
    task_dir = _task_dir(task_id)

    template_path = os.path.join(task_dir, "template.pptx")
    if not os.path.exists(template_path):
        raise ApiError("请先上传 PPT 模板。")

    output_path = os.path.join(task_dir, "output.pptx")
    if os.path.exists(output_path):
        os.remove(output_path)

    names = data.get("names", [])
    if not isinstance(names, list) or not names:
        raise ApiError("名单为空，请先导入或输入姓名。")

    settings = {
        "template_slide_index": data.get("template_slide_index", 0),
        "selected_box_indices": data.get("selected_box_indices", []),
        "repeat_count": data.get("repeat_count", 1),
        "fill_order": data.get("fill_order", "t2b"),
        "names_per_page": data.get("names_per_page"),
        "keep_blank_last_page": data.get("keep_blank_last_page", True),
        "treat_existing_as_placeholder": data.get("treat_existing_as_placeholder", True),
        "name_range": data.get("name_range"),
        "trim_whitespace": data.get("trim_whitespace", True),
    }

    try:
        stats = ppt_processor.generate(template_path, names, settings, output_path)
    except PptError as err:
        raise ApiError(str(err))
    except Exception as err:
        app.logger.exception("generate failed")
        raise ApiError("生成失败：" + str(err))

    # 文件名：原文件名_台卡生成结果_日期时间.pptx
    import datetime

    base = os.path.splitext(_original_template_name(task_dir))[0] or "台卡"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_display = f"{base}_台卡生成结果_{stamp}.pptx"

    return jsonify(ok=True, stats=stats, download_url=f"/api/download/{task_id}",
                   output_filename=output_display)


# ---------------------------------------------------------------------------
# 下载
# ---------------------------------------------------------------------------
@app.route("/api/download/<task_id>", methods=["GET"])
def download(task_id):
    task_dir = _task_dir(task_id)
    output_path = os.path.join(task_dir, "output.pptx")
    if not os.path.exists(output_path):
        raise ApiError("生成结果不存在，请重新生成。", 404)

    import datetime

    base = os.path.splitext(_original_template_name(task_dir))[0] or "台卡"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_from_directory(
        task_dir,
        "output.pptx",
        as_attachment=True,
        download_name=f"{base}_台卡生成结果_{stamp}.pptx",
    )


if __name__ == "__main__":
    _sweep_old_tasks()
    app.run(host="127.0.0.1", port=5000, debug=False)
