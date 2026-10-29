# -*- coding: utf-8 -*-
"""Flask 后端端到端流程测试：上传 → 分析 → 导入姓名 → 生成 → 下载。"""

from pptx import Presentation
from pptx.util import Inches

import app as app_module


def _make_template(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    box = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(4), Inches(1.5))
    box.text_frame.text = "姓名"
    p = tmp_path / "t.pptx"
    prs.save(str(p))
    return p


def test_full_flow(tmp_path):
    client = app_module.app.test_client()

    # 1) 新任务
    r = client.post("/api/new_task", json={})
    data = r.get_json()
    assert data["ok"]
    task_id = data["task_id"]

    # 2) 上传模板
    tpl = _make_template(tmp_path)
    with open(tpl, "rb") as f:
        r = client.post(
            "/api/upload_ppt",
            data={"task_id": task_id, "file": (f, "台卡模板.pptx")},
            content_type="multipart/form-data",
        )
    data = r.get_json()
    assert data["ok"], data
    assert data["analysis"]["slide_count"] == 1
    assert data["analysis"]["slides"][0]["text_shapes"][0]["text"] == "姓名"

    # 3) 粘贴导入姓名
    r = client.post("/api/import_names", data={"task_id": task_id, "text": "甲\n乙\n"})
    data = r.get_json()
    assert data["ok"]
    assert data["names"] == ["甲", "乙"]

    # 4) 生成
    payload = {
        "task_id": task_id,
        "names": ["甲", "乙", "丙"],
        "template_slide_index": 0,
        "selected_box_indices": [0],
        "repeat_count": 1,
        "fill_order": "t2b",
        "keep_blank_last_page": True,
        "treat_existing_as_placeholder": True,
    }
    r = client.post("/api/generate", json=payload)
    data = r.get_json()
    assert data["ok"], data
    assert data["stats"]["total_names"] == 3
    assert data["stats"]["total_pages"] == 3

    # 5) 下载
    r = client.get(data["download_url"])
    assert r.status_code == 200
    assert len(r.data) > 0
    assert r.data[:2] == b"PK"  # zip / pptx 魔数


def test_generate_without_template_fails(tmp_path):
    client = app_module.app.test_client()
    r = client.post("/api/new_task", json={})
    task_id = r.get_json()["task_id"]
    r = client.post(
        "/api/generate",
        json={"task_id": task_id, "names": ["甲"], "selected_box_indices": [0]},
    )
    assert r.status_code == 400
    assert r.get_json()["ok"] is False
