# -*- coding: utf-8 -*-
"""核心逻辑测试：重复次数、排列顺序、空白框、多页生成、输出可重新读取。"""

from pptx import Presentation
from pptx.util import Inches

import ppt_processor
from ppt_processor import PptError


# ---------------------------------------------------------------------------
# 测试用模板构造
# ---------------------------------------------------------------------------
def add_box(slide, left, top, width, height, text):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.text_frame.text = text
    return box


def make_two_box_template(path):
    """两个姓名框（上/下） + 一个固定文字 + 一个空白框。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    add_box(slide, 2, 2, 4, 1.2, "姓名A")   # index 0（上）
    add_box(slide, 2, 4, 4, 1.2, "姓名B")   # index 1（下）
    add_box(slide, 9, 1, 2, 0.8, "LOGO")    # index 2（固定）
    add_box(slide, 8, 5, 2, 0.8, "")        # index 3（空白）
    prs.save(path)
    return path


def make_grid_template(path):
    """2×2 姓名框网格，用于测试填写顺序。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(slide, 2, 1, 3, 1.0, "A")   # 左列上 index 0
    add_box(slide, 6, 1, 3, 1.0, "B")   # 右列上 index 1
    add_box(slide, 2, 4, 3, 1.0, "C")   # 左列下 index 2
    add_box(slide, 6, 4, 3, 1.0, "D")   # 右列下 index 3
    prs.save(path)
    return path


def read_all_texts(path):
    """重新打开输出文件，返回每一页的文本形状文字列表。"""
    prs = Presentation(path)
    result = []
    for slide in prs.slides:
        texts = [s.text_frame.text for s in ppt_processor.iter_text_shapes(slide)]
        result.append(texts)
    return result


def generate(path, names, selected, **kwargs):
    settings = {
        "template_slide_index": 0,
        "selected_box_indices": selected,
        "repeat_count": kwargs.get("repeat_count", 1),
        "fill_order": kwargs.get("fill_order", "t2b"),
        "names_per_page": kwargs.get("names_per_page"),
        "keep_blank_last_page": kwargs.get("keep_blank_last_page", True),
        "treat_existing_as_placeholder": True,
        "name_range": kwargs.get("name_range"),
        "trim_whitespace": True,
        "same_name_all_boxes": kwargs.get("same_name_all_boxes", False),
    }
    out = str(path).replace(".pptx", "_out.pptx")
    stats = ppt_processor.generate(str(path), names, settings, out)
    return out, stats


# ---------------------------------------------------------------------------
# 测试
# ---------------------------------------------------------------------------
def test_analyze_returns_text_shapes(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    info = ppt_processor.analyze_pptx(str(p))
    assert info["slide_count"] == 1
    shapes = info["slides"][0]["text_shapes"]
    assert len(shapes) == 4
    assert shapes[0]["text"] == "姓名A"
    assert shapes[3]["is_empty"] is True


def test_duplicate_slide_preserves_shapes(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    prs = Presentation(str(p))
    src = prs.slides[0]
    dup = ppt_processor.duplicate_slide(prs, src)
    src_texts = [s.text_frame.text for s in ppt_processor.iter_text_shapes(src)]
    dup_texts = [s.text_frame.text for s in ppt_processor.iter_text_shapes(dup)]
    assert src_texts == dup_texts


def test_generate_repeat_and_multipage(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["陈飘", "张三", "李四"]
    out, stats = generate(p, names, [0, 1], repeat_count=2, names_per_page=2)

    assert stats["total_names"] == 6
    assert stats["total_pages"] == 3

    pages = read_all_texts(out)
    assert len(pages) == 3
    # 每页两个姓名框顺序：姓名A(上)、姓名B(下)
    assert pages[0][0] == "陈飘" and pages[0][1] == "陈飘"
    assert pages[1][0] == "张三" and pages[1][1] == "张三"
    assert pages[2][0] == "李四" and pages[2][1] == "李四"


def test_generate_fill_order_l2r(tmp_path):
    p = make_grid_template(tmp_path / "t.pptx")
    names = ["N1", "N2", "N3", "N4"]
    out, stats = generate(p, names, [0, 1, 2, 3], fill_order="l2r", names_per_page=4)
    assert stats["total_pages"] == 1
    page = read_all_texts(out)[0]
    # 从左到右（先列）：index0=N1, index2=N2, index1=N3, index3=N4
    assert page[0] == "N1"
    assert page[1] == "N3"
    assert page[2] == "N2"
    assert page[3] == "N4"


def test_generate_fill_order_t2b(tmp_path):
    p = make_grid_template(tmp_path / "t.pptx")
    names = ["N1", "N2", "N3", "N4"]
    out, _ = generate(p, names, [0, 1, 2, 3], fill_order="t2b", names_per_page=4)
    page = read_all_texts(out)[0]
    # 从上到下（先行）：index0=N1, index1=N2, index2=N3, index3=N4
    assert page[0] == "N1"
    assert page[1] == "N2"
    assert page[2] == "N3"
    assert page[3] == "N4"


def test_empty_box_can_be_filled(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["王五"]
    out, _ = generate(p, names, [3], repeat_count=1, names_per_page=1)
    page = read_all_texts(out)[0]
    assert page[3] == "王五"  # 空白框被填入姓名


def test_name_range(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["A", "B", "C", "D"]
    out, stats = generate(p, names, [0, 1], repeat_count=1, names_per_page=2,
                          name_range="2-3")
    assert stats["total_names"] == 2
    assert stats["total_pages"] == 1
    page = read_all_texts(out)[0]
    assert page[0] == "B"
    assert page[1] == "C"


def test_keep_blank_last_page_true(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["A", "B", "C"]
    out, _ = generate(p, names, [0, 1], repeat_count=1, names_per_page=2,
                      keep_blank_last_page=True)
    pages = read_all_texts(out)
    assert len(pages) == 2
    assert pages[1][0] == "C"
    assert pages[1][1] == ""  # 未使用位置被清空


def test_keep_blank_last_page_false(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["A", "B", "C"]
    out, _ = generate(p, names, [0, 1], repeat_count=1, names_per_page=2,
                      keep_blank_last_page=False)
    pages = read_all_texts(out)
    assert pages[1][0] == "C"
    assert pages[1][1] == "姓名B"  # 未使用位置保留原文字


def test_empty_names_raises(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    try:
        generate(p, [], [0, 1])
    except PptError:
        return
    raise AssertionError("空名单应当抛出 PptError")


def test_no_selection_raises(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    try:
        generate(p, ["A"], [])
    except PptError:
        return
    raise AssertionError("未选择姓名框应当抛出 PptError")


def test_same_name_all_boxes(tmp_path):
    """折叠式台卡：同一姓名填入所有选中框（每页 1 个姓名）。"""
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["陈飘", "张三", "李四"]
    out, stats = generate(p, names, [0, 1], repeat_count=1, same_name_all_boxes=True)
    assert stats["total_pages"] == 3
    assert stats["names_per_page"] == 1
    pages = read_all_texts(out)
    assert len(pages) == 3
    # 每一页的两个姓名框都是同一个姓名
    assert pages[0][0] == "陈飘" and pages[0][1] == "陈飘"
    assert pages[1][0] == "张三" and pages[1][1] == "张三"
    assert pages[2][0] == "李四" and pages[2][1] == "李四"


def test_output_reopens_and_counts_match(tmp_path):
    p = make_two_box_template(tmp_path / "t.pptx")
    names = ["甲", "乙", "丙", "丁", "戊"]
    out, stats = generate(p, names, [0, 1], repeat_count=2, names_per_page=2)
    # 10 个姓名 / 每页 2 个 = 5 页
    assert stats["total_names"] == 10
    assert stats["total_pages"] == 5

    prs = Presentation(out)
    assert len(prs.slides) == 5
    # 每页文本形状数量与原模板一致（4 个）
    for slide in prs.slides:
        assert len(ppt_processor.iter_text_shapes(slide)) == 4
